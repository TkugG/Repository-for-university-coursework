import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ── Projector-Safe Palette (SoftBank Light Theme) ──────────────────────────────
C_BG         = RGBColor(255, 255, 255)  # Pure White
C_LINE       = RGBColor(226, 232, 240)  # Slate 200 (Hairline only)
C_NAVY       = RGBColor(30,  58,  138)  # #1E3A8A Deep Navy
C_BLUE       = RGBColor(37,  99,  235)  # #2563EB Vivid Blue
C_TEXT_MAIN  = RGBColor(15,  23,  42)   # #0F172A Slate 900
C_TEXT_MUTED = RGBColor(100, 116, 139)  # #64748B Slate 500
FONT_FAMILY  = "Prompt"

# ── Grid Constants (13.333" x 7.500" — 0.90" margin) ─────────────────────────
MARGIN_L  = Inches(0.90)
CONTENT_W = Inches(11.533)

COL2_W  = Inches(5.466);  COL2_X1 = Inches(0.90);  COL2_X2 = Inches(6.966)
COL3_W  = Inches(3.544);  COL3_X1 = Inches(0.90);  COL3_X2 = Inches(4.894);  COL3_X3 = Inches(8.888)

# ── Helpers (100% Isolated Textboxes & Zero Padding) ──────────────────────────

def apply_canvas(slide):
    """Sets pure white background. MUST be the first call for every slide."""
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_BG
    bg.line.fill.background()
    return bg

def add_textbox(slide, left, top, width, height, text,
                font_size=Pt(14), bold=False, color=C_TEXT_MAIN,
                align=PP_ALIGN.LEFT, word_wrap=True):
    """Canva-safe isolated textbox — 1 font size, 0 internal padding."""
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = word_wrap
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = text
    p.font.name      = FONT_FAMILY
    p.font.size      = font_size
    p.font.bold      = bold
    p.font.color.rgb = color
    p.alignment      = align
    return tb

def add_clean_header(slide, category_text, title_text):
    """Standard header: Category Tag + Slide Title + Hairline Divider."""
    add_textbox(slide, MARGIN_L, Inches(0.45), CONTENT_W, Inches(0.35),
                category_text.upper(), font_size=Pt(11), bold=True, color=C_BLUE)
    add_textbox(slide, MARGIN_L, Inches(0.85), CONTENT_W, Inches(0.80),
                title_text, font_size=Pt(24), bold=True, color=C_TEXT_MAIN)
    add_hairline(slide, MARGIN_L, Inches(1.85), CONTENT_W)

def add_hairline(slide, left, top, width, vertical=False):
    """1px hairline divider — Slate 200."""
    size = (Inches(0.015), width) if vertical else (width, Inches(0.015))
    div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, *size)
    div.fill.solid()
    div.fill.fore_color.rgb = C_LINE
    div.line.fill.background()
    return div

# ── Main Presentation Builder ─────────────────────────────────────────────────

def build_presentation(output_path):
    prs = Presentation()
    prs.slide_width  = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 1: หน้าปก (COVER)
    # ─────────────────────────────────────────────────────────────────────────
    s1 = prs.slides.add_slide(blank)
    apply_canvas(s1)
    
    add_textbox(s1, MARGIN_L, Inches(1.80), CONTENT_W, Inches(0.40),
                "PHP OOP & OOAD PROJECT DEFENSE", font_size=Pt(12), bold=True, color=C_BLUE)
    
    add_textbox(s1, MARGIN_L, Inches(2.35), CONTENT_W, Inches(1.20),
                "ระบบบันทึกและจัดการสิ่งที่ต้องทำ", font_size=Pt(48), bold=True, color=C_NAVY)
    
    add_textbox(s1, MARGIN_L, Inches(3.65), CONTENT_W, Inches(0.60),
                "การออกแบบสถาปัตยกรรมส่วนหลังบ้านด้วย PHP OOP และ REST API", font_size=Pt(18), color=C_TEXT_MUTED)
    
    add_hairline(s1, MARGIN_L, Inches(4.55), CONTENT_W)

    # Metadata 2 คอลัมน์
    add_textbox(s1, MARGIN_L, Inches(4.90), COL2_W, Inches(0.35),
                "ผู้รับผิดชอบส่วนหลังบ้าน (Backend Developer)", font_size=Pt(12), bold=True, color=C_BLUE)
    add_textbox(s1, MARGIN_L, Inches(5.30), COL2_W, Inches(0.40),
                "นักศึกษาผู้พัฒนาส่วนฐานข้อมูลและคลาสเชิงวัตถุ", font_size=Pt(14), color=C_TEXT_MAIN)

    add_textbox(s1, COL2_X2, Inches(4.90), COL2_W, Inches(0.35),
                "สถาปัตยกรรมระบบ (Architecture)", font_size=Pt(12), bold=True, color=C_BLUE)
    add_textbox(s1, COL2_X2, Inches(5.30), COL2_W, Inches(0.40),
                "Decoupled MVC Pattern / MySQL / PDO", font_size=Pt(14), color=C_TEXT_MAIN)

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 2: สถาปัตยกรรมระบบ (SYSTEM ARCHITECTURE)
    # ─────────────────────────────────────────────────────────────────────────
    s2 = prs.slides.add_slide(blank)
    apply_canvas(s2)
    add_clean_header(s2, "SYSTEM ARCHITECTURE", "สถาปัตยกรรมระบบแบบ Decoupled MVC")

    # Col 1: View
    add_textbox(s2, COL3_X1, Inches(2.20), COL3_W, Inches(0.35), "01", font_size=Pt(20), bold=True, color=C_BLUE)
    add_textbox(s2, COL3_X1, Inches(2.65), COL3_W, Inches(0.40), "View (หน้าบ้าน)", font_size=Pt(18), bold=True, color=C_NAVY)
    add_textbox(s2, COL3_X1, Inches(3.20), COL3_W, Inches(2.20),
                "• พัฒนาด้วย HTML5, CSS3 และ JavaScript\n"
                "• แสดงผลแบบ Single Page Application\n"
                "• ติดต่อหลังบ้านผ่านคำขอแบบ Fetch API\n"
                "• ไม่มีการเขียนโค้ดต่อฐานข้อมูลที่ฝั่งผู้ใช้",
                font_size=Pt(13), color=C_TEXT_MAIN)

    # Col 2: Controller
    add_textbox(s2, COL3_X2, Inches(2.20), COL3_W, Inches(0.35), "02", font_size=Pt(20), bold=True, color=C_BLUE)
    add_textbox(s2, COL3_X2, Inches(2.65), COL3_W, Inches(0.40), "Controller (ตัวกลาง API)", font_size=Pt(18), bold=True, color=C_NAVY)
    add_textbox(s2, COL3_X2, Inches(3.20), COL3_W, Inches(2.20),
                "• อยู่ในโฟลเดอร์ api/ ทำหน้าที่เป็น REST API\n"
                "• รับคำสั่ง GET, POST, DELETE จากหน้าเว็บ\n"
                "• ตรวจสอบความถูกต้องของข้อมูลที่ส่งมา\n"
                "• ตอบกลับผลลัพธ์ในรูปแบบมาตรฐาน JSON",
                font_size=Pt(13), color=C_TEXT_MAIN)

    # Col 3: Model
    add_textbox(s2, COL3_X3, Inches(2.20), COL3_W, Inches(0.35), "03", font_size=Pt(20), bold=True, color=C_BLUE)
    add_textbox(s2, COL3_X3, Inches(2.65), COL3_W, Inches(0.40), "Model (ตรรกะเชิงวัตถุ)", font_size=Pt(18), bold=True, color=C_NAVY)
    add_textbox(s2, COL3_X3, Inches(3.20), COL3_W, Inches(2.20),
                "• คลาส PHP OOP แท้ๆ ในโฟลเดอร์ models/\n"
                "• รับผิดชอบการคำนวณและคำสั่ง SQL ทั้งหมด\n"
                "• เชื่อมต่อฐานข้อมูล MySQL ผ่านทาง PDO\n"
                "• แยกหน้าที่ชัดเจน ไม่ปะปนกับส่วนแสดงผล",
                font_size=Pt(13), color=C_TEXT_MAIN)

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 3: โครงสร้างคลาส (DOMAIN MODELING & UML)
    # ─────────────────────────────────────────────────────────────────────────
    s3 = prs.slides.add_slide(blank)
    apply_canvas(s3)
    add_clean_header(s3, "DOMAIN MODELING", "โครงสร้างคลาสตามหลักการโปรแกรมเชิงวัตถุ")

    # Col 1: Database Class
    add_textbox(s3, COL3_X1, Inches(2.20), COL3_W, Inches(0.35), "CONNECTION", font_size=Pt(11), bold=True, color=C_BLUE)
    add_textbox(s3, COL3_X1, Inches(2.60), COL3_W, Inches(0.40), "Class Database", font_size=Pt(18), bold=True, color=C_NAVY)
    add_textbox(s3, COL3_X1, Inches(3.10), COL3_W, Inches(2.60),
                "• Attributes:\n"
                "  - host, db_name, username, password (private)\n"
                "  - conn (PDO connection)\n"
                "• Methods:\n"
                "  + connect(): PDO\n\n"
                "จัดการการเปิดจุดเชื่อมต่อ MySQL เพียงจุดเดียว",
                font_size=Pt(13), color=C_TEXT_MAIN)

    # Col 2: Note Class
    add_textbox(s3, COL3_X2, Inches(2.20), COL3_W, Inches(0.35), "CORE ENTITY", font_size=Pt(11), bold=True, color=C_BLUE)
    add_textbox(s3, COL3_X2, Inches(2.60), COL3_W, Inches(0.40), "Class Note", font_size=Pt(18), bold=True, color=C_NAVY)
    add_textbox(s3, COL3_X2, Inches(3.10), COL3_W, Inches(2.60),
                "• Attributes:\n"
                "  - conn, table (private)\n"
                "  + id, title, type, content, color, reminder_at\n"
                "• Methods:\n"
                "  + __construct(db)\n"
                "  + readAll(sort, category, keyword): array\n"
                "  + create(data): int / update(id, data): bool\n"
                "  + delete(id): bool / getReminders(): array",
                font_size=Pt(13), color=C_TEXT_MAIN)

    # Col 3: TodoItem Class
    add_textbox(s3, COL3_X3, Inches(2.20), COL3_W, Inches(0.35), "SUB ENTITY", font_size=Pt(11), bold=True, color=C_BLUE)
    add_textbox(s3, COL3_X3, Inches(2.60), COL3_W, Inches(0.40), "Class TodoItem", font_size=Pt(18), bold=True, color=C_NAVY)
    add_textbox(s3, COL3_X3, Inches(3.10), COL3_W, Inches(2.60),
                "• Attributes:\n"
                "  - conn, table (private)\n"
                "• Methods:\n"
                "  + __construct(db)\n"
                "  + toggle(id, is_done): bool\n"
                "  + create(note_id, text): int\n"
                "  + delete(id): bool\n\n"
                "จัดการรายการย่อย ความสัมพันธ์แบบ 1 ต่อกลุ่ม",
                font_size=Pt(13), color=C_TEXT_MAIN)

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 4: หัวใจหลัก OOP (CORE OOP PRINCIPLES)
    # ─────────────────────────────────────────────────────────────────────────
    s4 = prs.slides.add_slide(blank)
    apply_canvas(s4)
    add_clean_header(s4, "OBJECT-ORIENTED PRINCIPLES", "การประยุกต์ใช้หลักการเชิงวัตถุในโค้ด")

    # Col 1: Encapsulation
    add_textbox(s4, COL2_X1, Inches(2.20), COL2_W, Inches(0.35), "PRINCIPLE 01", font_size=Pt(11), bold=True, color=C_BLUE)
    add_textbox(s4, COL2_X1, Inches(2.60), COL2_W, Inches(0.40), "Encapsulation (การห่อหุ้ม)", font_size=Pt(18), bold=True, color=C_NAVY)
    add_textbox(s4, COL2_X1, Inches(3.15), COL2_W, Inches(2.60),
                "• ซ่อนข้อมูลการเชื่อมต่อฐานข้อมูลเป็น private ใน Database\n"
                "• ตัวแปร $conn และ $table ใน Model ทุกตัวถูกล็อกเป็น private\n"
                "• คลาสภายนอกไม่สามารถเข้าถึงตัวแปรสำคัญได้โดยตรง\n"
                "• เข้าถึงข้อมูลผ่านทาง Method ที่กำหนดไว้เท่านั้น\n"
                "• ป้องกันข้อผิดพลาดจากการแก้ไขค่าตัวแปรโดยพลการ",
                font_size=Pt(13), color=C_TEXT_MAIN)

    # Col 2: Dependency Injection
    add_textbox(s4, COL2_X2, Inches(2.20), COL2_W, Inches(0.35), "PRINCIPLE 02", font_size=Pt(11), bold=True, color=C_BLUE)
    add_textbox(s4, COL2_X2, Inches(2.60), COL2_W, Inches(0.40), "Dependency Injection (การส่งต่ออ็อบเจกต์)", font_size=Pt(18), bold=True, color=C_NAVY)
    add_textbox(s4, COL2_X2, Inches(3.15), COL2_W, Inches(2.60),
                "• คลาส Note และ TodoItem รับการเชื่อมต่อผ่าน __construct($db)\n"
                "• ไม่สร้างอ็อบเจกต์การเชื่อมต่อซ้ำซ้อนภายในคลาส\n"
                "• ระบบทั้งหมดใช้งาน Connection เดียวกันอย่างมีประสิทธิภาพ\n"
                "• ลดการใช้หน่วยความจำและทรัพยากรของเซิร์ฟเวอร์\n"
                "• เอื้อต่อการทดสอบการทำงานของแต่ละคลาสได้อย่างอิสระ",
                font_size=Pt(13), color=C_TEXT_MAIN)

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 5: การแยกหน้าที่และการประมวลผล (RESPONSIBILITY & WORKFLOW)
    # ─────────────────────────────────────────────────────────────────────────
    s5 = prs.slides.add_slide(blank)
    apply_canvas(s5)
    add_clean_header(s5, "RESPONSIBILITY & WORKFLOW", "การแยกขอบเขตหน้าที่และการประมวลผลข้อมูล")

    # Col 1: Single Responsibility
    add_textbox(s5, COL2_X1, Inches(2.20), COL2_W, Inches(0.35), "ARCHITECTURE", font_size=Pt(11), bold=True, color=C_BLUE)
    add_textbox(s5, COL2_X1, Inches(2.60), COL2_W, Inches(0.40), "Single Responsibility Principle", font_size=Pt(18), bold=True, color=C_NAVY)
    add_textbox(s5, COL2_X1, Inches(3.15), COL2_W, Inches(2.60),
                "• แยกคลาส TodoItem ออกมาจากคลาส Note อย่างเด็ดขาด\n"
                "• Note ดูแลหัวข้อ สี หมวดหมู่ และการค้นหาภาพรวม\n"
                "• TodoItem ดูแลเฉพาะการติ๊กถูกขีดฆ่า และรายการย่อย\n"
                "• เมธอด toggle() ทำงานแบบเฉพาะจุด ไม่ต้องโหลดโน๊ตทั้งใบ\n"
                "• ทำให้โค้ดอ่านง่าย แก้ไขข้อบกพร่องได้รวดเร็ว",
                font_size=Pt(13), color=C_TEXT_MAIN)

    # Col 2: Action Routing
    add_textbox(s5, COL2_X2, Inches(2.20), COL2_W, Inches(0.35), "API CONTROLLER", font_size=Pt(11), bold=True, color=C_BLUE)
    add_textbox(s5, COL2_X2, Inches(2.60), COL2_W, Inches(0.40), "RESTful Action Routing", font_size=Pt(18), bold=True, color=C_NAVY)
    add_textbox(s5, COL2_X2, Inches(3.15), COL2_W, Inches(2.60),
                "• ไฟล์ api/notes.php รองรับคำขออย่างยืดหยุ่น\n"
                "• GET: รองรับการจัดเรียง การค้นหา และการดึงโน๊ตแจ้งเตือน\n"
                "• POST: รองรับทั้งการสร้างใหม่ และอัปเดตผ่าน ?action=update\n"
                "• Interactive Action: รองรับ ?action=toggle_todo จากหน้าเว็บ\n"
                "• ส่งออก Header Content-Type: application/json มาตรฐานสากล",
                font_size=Pt(13), color=C_TEXT_MAIN)

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 6: ความปลอดภัยและความถูกต้องของข้อมูล (SECURITY & INTEGRITY)
    # ─────────────────────────────────────────────────────────────────────────
    s6 = prs.slides.add_slide(blank)
    apply_canvas(s6)
    add_clean_header(s6, "DATA INTEGRITY & SECURITY", "ความปลอดภัยและความถูกต้องของข้อมูล")

    # Col 1: SQL Injection
    add_textbox(s6, COL3_X1, Inches(2.20), COL3_W, Inches(0.35), "DEFENSE 01", font_size=Pt(11), bold=True, color=C_BLUE)
    add_textbox(s6, COL3_X1, Inches(2.60), COL3_W, Inches(0.40), "PDO Prepared Statements", font_size=Pt(18), bold=True, color=C_NAVY)
    add_textbox(s6, COL3_X1, Inches(3.15), COL3_W, Inches(2.40),
                "• ใช้ bindValue() ในทุกคำสั่ง\n"
                "• ป้องกันการโจมตีแบบ SQL Injection\n"
                "• แยกคำสั่ง SQL ออกจากข้อมูลนำเข้า\n"
                "• ปลอดภัยกว่าการต่อข้อความแบบดั้งเดิม",
                font_size=Pt(13), color=C_TEXT_MAIN)

    # Col 2: File Upload
    add_textbox(s6, COL3_X2, Inches(2.20), COL3_W, Inches(0.35), "DEFENSE 02", font_size=Pt(11), bold=True, color=C_BLUE)
    add_textbox(s6, COL3_X2, Inches(2.60), COL3_W, Inches(0.40), "Unique Asset Naming", font_size=Pt(18), bold=True, color=C_NAVY)
    add_textbox(s6, COL3_X2, Inches(3.15), COL3_W, Inches(2.40),
                "• ใช้คำสั่ง uniqid('note_', true)\n"
                "• สุ่มสร้างชื่อไฟล์ตามเวลาในระดับไมโครวินาที\n"
                "• ป้องกันการอัปโหลดไฟล์ชื่อซ้ำทับซ้อนกัน\n"
                "• แยกเก็บไฟล์ในโฟลเดอร์ uploads/",
                font_size=Pt(13), color=C_TEXT_MAIN)

    # Col 3: Cascade Delete
    add_textbox(s6, COL3_X3, Inches(2.20), COL3_W, Inches(0.35), "DEFENSE 03", font_size=Pt(11), bold=True, color=C_BLUE)
    add_textbox(s6, COL3_X3, Inches(2.60), COL3_W, Inches(0.40), "ON DELETE CASCADE", font_size=Pt(18), bold=True, color=C_NAVY)
    add_textbox(s6, COL3_X3, Inches(3.15), COL3_W, Inches(2.40),
                "• กำหนด Foreign Key ระดับฐานข้อมูล\n"
                "• เมื่อลบโน๊ตแม่ ข้อย่อยจะถูกลบตามทันที\n"
                "• ไม่ต้องเขียนคำสั่งลบหลายรอบในโค้ด\n"
                "• ป้องกันปัญหาข้อมูลลูกกำพร้าตกค้าง",
                font_size=Pt(13), color=C_TEXT_MAIN)

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 7: บทสรุปและคำถาม (CONCLUSION & Q&A)
    # ─────────────────────────────────────────────────────────────────────────
    s7 = prs.slides.add_slide(blank)
    apply_canvas(s7)
    add_clean_header(s7, "SUMMARY & DEFENSE", "สรุปผลการพัฒนาระบบหลังบ้าน")

    # Col 1: สรุป
    add_textbox(s7, COL2_X1, Inches(2.20), COL2_W, Inches(0.35), "ACHIEVEMENT", font_size=Pt(11), bold=True, color=C_BLUE)
    add_textbox(s7, COL2_X1, Inches(2.60), COL2_W, Inches(0.40), "ความสำเร็จของการพัฒนาระบบ", font_size=Pt(18), bold=True, color=C_NAVY)
    add_textbox(s7, COL2_X1, Inches(3.15), COL2_W, Inches(2.60),
                "• ตอบโจทย์หลักการ OOP ครบถ้วน (Encapsulation, DI, SRP)\n"
                "• สถาปัตยกรรม Decoupled MVC รองรับการทำงานคู่ขนานได้อย่างราบรื่น\n"
                "• ฐานข้อมูลมีความสมบูรณ์ รองรับทั้งการค้นหา กรองหมวด และการแจ้งเตือน\n"
                "• โค้ดมีระเบียบ พร้อมสำหรับการนำไปขยายผลต่อในอนาคต",
                font_size=Pt(13), color=C_TEXT_MAIN)

    # Col 2: Q&A
    add_textbox(s7, COL2_X2, Inches(2.20), COL2_W, Inches(0.35), "READY FOR DEFENSE", font_size=Pt(11), bold=True, color=C_BLUE)
    add_textbox(s7, COL2_X2, Inches(2.60), COL2_W, Inches(0.40), "ขอขอบคุณคณะกรรมการและอาจารย์", font_size=Pt(18), bold=True, color=C_NAVY)
    add_textbox(s7, COL2_X2, Inches(3.15), COL2_W, Inches(2.60),
                "ระบบหลังบ้านพร้อมสำหรับการสาธิตการทำงานจริง\n\n"
                "ยินดีตอบทุกข้อซักถามและรับคำแนะนำเพื่อการพัฒนาต่อไปครับ",
                font_size=Pt(14), color=C_TEXT_MUTED)

    # บันทึกไฟล์
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    out_dir = r"c:\xampp\htdocs\ProjectOOP\slid"
    os.makedirs(out_dir, exist_ok=True)
    out_file = os.path.join(out_dir, "backend_presentation.pptx")
    build_presentation(out_file)
